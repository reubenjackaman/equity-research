"""
Checks that the exported DCF workbook's formulas give exactly the same answers as the
Python model, by recalculating the .xlsx with the `formulas` library (an Excel engine
written in Python).

One-time setup:   .venv/bin/pip install formulas
Run from backend: .venv/bin/python -m tests.check_dcf_excel
"""
import tempfile
from pathlib import Path

import formulas
import openpyxl
from dotenv import load_dotenv
from openpyxl.utils import get_column_letter as L

load_dotenv(".env")
from exports import dcf_excel  # noqa: E402
from finance import dcf  # noqa: E402

CASES = {
    "AAPL defaults": ("AAPL", {}),
    "AAPL custom, 8 years": ("AAPL", {"years": 8, "growth_start": 0.07, "growth_end": 0.03, "margin_target": 0.34, "erp": 0.045}),
    "AAPL WACC override": ("AAPL", {"wacc_override": 0.085}),
    "SHOP.TO (USD reporting, CAD price)": ("SHOP.TO", {}),
}


def close(x, y):
    if x is None:
        return y in (None, "N/A")
    return isinstance(y, (int, float)) and abs(x - y) <= max(1e-6, abs(x) * 1e-6)


bad = 0
with tempfile.TemporaryDirectory() as tmp:
    for name, (ticker, overrides) in CASES.items():
        m = dcf.model(ticker, overrides)
        path = Path(tmp) / "dcf.xlsx"
        path.write_bytes(dcf_excel.workbook(m))

        solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
        cells = {}
        for ref, value in solution.items():
            try:
                cells[ref.split("!")[-1].upper()] = value.value[0][0]
            except Exception:
                pass

        ws = openpyxl.load_workbook(path).active
        rows = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
        per_share_row = max(r for label, r in rows.items() if isinstance(label, str) and label.startswith("Value per share ("))
        first_forecast_col = 2 + len(m["history"])
        sens_row = rows["WACC ↓ / Terminal growth →"]

        excel_fcf = [cells.get(f"{L(first_forecast_col + i)}{rows['Unlevered free cash flow']}") for i in range(m["years"])]
        excel_grid = [[cells.get(f"{L(2 + j)}{sens_row + 1 + i}") for j in range(5)] for i in range(5)]
        checks = {
            "WACC": close(m["wacc"]["used"], cells.get("B25")),
            "FCFs": all(close(f["fcf"] / 1e6, x) for f, x in zip(m["forecast"], excel_fcf)),
            "EV": close(m["valuation"]["enterpriseValue"] / 1e6, cells.get(f"B{rows['Enterprise value']}")),
            "value/share": close(m["valuation"]["perShareTrading"], cells.get(f"B{per_share_row}")),
            "sensitivity": all(close(a, b) for ra, rb in zip(m["sensitivity"]["values"], excel_grid) for a, b in zip(ra, rb)),
        }
        bad += not all(checks.values())
        print("OK  " if all(checks.values()) else "DIFF", f"{name:36} value/share {m['valuation']['perShareTrading']:.2f}", checks)
print("mismatching workbooks:", bad)
