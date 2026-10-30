"""
DCF export to Excel with live formulas.

Layout follows the usual banking convention: blue cells are inputs, black cells are
formulas. Change any blue cell in Excel and the forecast, valuation and sensitivity
table all recalculate. The formulas mirror finance/dcf.py exactly.

The number of forecast years is fixed when exported (it decides how many columns exist);
export again from the app to change it.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

INPUT = Font(color="0000FF")
BOLD = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOP_BORDER = Border(top=Side(style="thin"))
PCT = "0.0%"
PCT2 = "0.00%"
MONEY = '#,##0;(#,##0);"–"'
PRICE = "#,##0.00"
M = 1e6


def workbook(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "DCF"
    a = result["assumptions"]
    v = result["valuation"]
    w = result["wacc"]
    hist, n = result["history"], result["years"]
    ccy, tccy = result["currency"], result["tradingCurrency"]

    def put(cell, value, fmt=None, font=None):
        ws[cell] = value
        if fmt:
            ws[cell].number_format = fmt
        if font:
            ws[cell].font = font
        return ws[cell]

    def header(row, text):
        for c in range(1, 3 + len(hist) + n):
            ws.cell(row=row, column=c).fill = HEADER_FILL
        put(f"A{row}", text, font=HEADER_FONT)

    put("A1", f"{result['name']} ({result['symbol']}) · Discounted cash flow model", font=Font(bold=True, size=14))
    put("A2", f"{ccy} millions except per-share figures. Blue = inputs you can change; black = formulas. "
              "For educational purposes only, not investment advice.", font=Font(italic=True, color="64748B"))

    # ---------- Assumptions ----------
    header(4, "Operating assumptions")
    rows = [
        (5, "Forecast years (fixed: re-export to change)", n, "0", None),
        (6, "Revenue growth: year 1", a["growth_start"], PCT, INPUT),
        (7, "Revenue growth: final year", a["growth_end"], PCT, INPUT),
        (8, "Target EBIT margin (reached in final year)", a["margin_target"], PCT, INPUT),
        (9, "Tax rate", a["tax_rate"], PCT, INPUT),
        (10, "D&A, % of revenue", a["da_pct"], PCT2, INPUT),
        (11, "Capex, % of revenue", a["capex_pct"], PCT2, INPUT),
        (12, "Net working capital, % of revenue", a["nwc_pct"], PCT2, INPUT),
        (13, "Terminal growth rate", a["terminal_growth"], PCT2, INPUT),
    ]
    header(15, "Discount rate (WACC)")
    rows += [
        (16, "Risk-free rate (10-year Treasury)", a["risk_free"], PCT2, INPUT),
        (17, "Beta", a["beta"], "0.00", INPUT),
        (18, "Equity risk premium", a["erp"], PCT2, INPUT),
        (19, "Cost of equity = Rf + Beta × ERP", "=B16+B17*B18", PCT2, None),
        (20, "Pre-tax cost of debt", a["cost_of_debt"], PCT2, INPUT),
        (21, "After-tax cost of debt", "=B20*(1-B9)", PCT2, None),
        (22, f"Market cap ({ccy} M)", (result["marketCap"] or 0) / M, MONEY, INPUT),
        (23, f"Total debt ({ccy} M)", v["debt"] / M if v else 0, MONEY, INPUT),
        (24, "Equity weight E/(D+E)", "=IF(B22+B23>0,B22/(B22+B23),1)", PCT, None),
    ]
    if a.get("wacc_override") is not None:
        rows.append((25, "WACC (your override)", a["wacc_override"], PCT2, INPUT))
        put("C25", f"CAPM build-up gives {w['computed']:.2%}. To use it, enter =B24*B19+(1-B24)*B21 in B25.",
            font=Font(italic=True, color="64748B"))
    else:
        rows.append((25, "WACC = E/V × Ke + D/V × Kd × (1 − t)", "=B24*B19+(1-B24)*B21", PCT2, BOLD))
    header(27, f"Balance sheet & shares (as of {result['balanceAsOf']})")
    rows += [
        (28, f"Cash + short-term investments ({ccy} M)", v["cash"] / M if v else 0, MONEY, INPUT),
        (29, "Diluted shares (M)", (v["shares"] if v else 0) / M, "#,##0.0", INPUT),
        (30, f"Current share price ({tccy})", result["price"], PRICE, INPUT),
        (31, f"FX rate: {tccy} per 1 {ccy}", result["fx"], "0.0000", INPUT),
    ]
    for r, label, value, fmt, font in rows:
        put(f"A{r}", label)
        put(f"B{r}", value, fmt, font)

    # ---------- Forecast ----------
    top = 33
    header(top, f"Free cash flow forecast ({ccy} M)")
    first_f = 2 + len(hist)              # first forecast column number
    last_h = col(first_f - 1)            # last historical column letter
    last_f = col(first_f + n - 1)        # last forecast column letter
    fcols = [col(first_f + t) for t in range(n)]

    R = {name: top + 1 + i for i, name in enumerate(
        ["year", "revenue", "growth", "margin", "ebit", "taxes", "nopat", "da", "capex", "nwc", "dnwc", "fcf", "df", "pv"])}
    labels = {
        "year": "Forecast year #", "revenue": "Revenue", "growth": "Revenue growth", "margin": "EBIT margin",
        "ebit": "EBIT (operating income)", "taxes": "Taxes on EBIT", "nopat": "NOPAT = EBIT − taxes",
        "da": "+ Depreciation & amortization", "capex": "− Capital expenditures", "nwc": "Net working capital",
        "dnwc": "− Increase in net working capital", "fcf": "Unlevered free cash flow",
        "df": "Discount factor = 1 ÷ (1 + WACC)^t", "pv": "Present value of FCF",
    }
    for key, r in R.items():
        put(f"A{r}", labels[key], font=BOLD if key in ("fcf", "revenue") else None)

    for i, h in enumerate(hist):
        c = col(2 + i)
        put(f"{c}{top}", f"{h['label']} (actual)", font=HEADER_FONT)
        put(f"{c}{R['revenue']}", h["revenue"] / M, MONEY)
        put(f"{c}{R['growth']}", h["growth"], PCT)
        put(f"{c}{R['ebit']}", h["ebit"] / M, MONEY)
        put(f"{c}{R['margin']}", f"={c}{R['ebit']}/{c}{R['revenue']}", PCT)
        put(f"{c}{R['da']}", h["da"] / M if h["da"] is not None else "N/A", MONEY)
        put(f"{c}{R['capex']}", h["capex"] / M if h["capex"] is not None else "N/A", MONEY)
        if h["nwc"] is not None:
            put(f"{c}{R['nwc']}", h["nwc"] / M, MONEY)
        elif i == len(hist) - 1:  # same fallback as the Python model: assumed % of revenue
            put(f"{c}{R['nwc']}", f"={c}{R['revenue']}*$B$12", MONEY)
        else:
            put(f"{c}{R['nwc']}", "N/A")

    for t, c in enumerate(fcols, start=1):
        p = col(first_f + t - 2)
        put(f"{c}{top}", f"Year {t}", font=HEADER_FONT)
        put(f"{c}{R['year']}", t, "0")
        put(f"{c}{R['growth']}", f"=$B$6+($B$7-$B$6)*({c}{R['year']}-1)/($B$5-1)", PCT)
        put(f"{c}{R['revenue']}", f"={p}{R['revenue']}*(1+{c}{R['growth']})", MONEY)
        put(f"{c}{R['margin']}", f"=${last_h}${R['margin']}+($B$8-${last_h}${R['margin']})*{c}{R['year']}/$B$5", PCT)
        put(f"{c}{R['ebit']}", f"={c}{R['revenue']}*{c}{R['margin']}", MONEY)
        put(f"{c}{R['taxes']}", f"={c}{R['ebit']}*$B$9", MONEY)
        put(f"{c}{R['nopat']}", f"={c}{R['ebit']}-{c}{R['taxes']}", MONEY)
        put(f"{c}{R['da']}", f"={c}{R['revenue']}*$B$10", MONEY)
        put(f"{c}{R['capex']}", f"={c}{R['revenue']}*$B$11", MONEY)
        put(f"{c}{R['nwc']}", f"={c}{R['revenue']}*$B$12", MONEY)
        put(f"{c}{R['dnwc']}", f"={c}{R['nwc']}-{p}{R['nwc']}", MONEY)
        put(f"{c}{R['fcf']}", f"={c}{R['nopat']}+{c}{R['da']}-{c}{R['capex']}-{c}{R['dnwc']}", MONEY, BOLD)
        put(f"{c}{R['df']}", f"=1/(1+$B$25)^{c}{R['year']}", "0.0000")
        put(f"{c}{R['pv']}", f"={c}{R['fcf']}*{c}{R['df']}", MONEY)
    for c in range(1, first_f + n):
        ws.cell(row=R["fcf"], column=c).border = TOP_BORDER

    # ---------- Valuation ----------
    vt = R["pv"] + 2
    header(vt, "Valuation")
    fcf_range = f"{fcols[0]}{R['fcf']}:{last_f}{R['fcf']}"
    val = [
        ("Sum of PV of free cash flows", f"=SUM({fcols[0]}{R['pv']}:{last_f}{R['pv']})", MONEY),
        ("Terminal value = FCF_N × (1 + g) ÷ (WACC − g)",
         f'=IF($B$25<=$B$13,"N/A",{last_f}{R["fcf"]}*(1+$B$13)/($B$25-$B$13))', MONEY),
        ("PV of terminal value", f'=IFERROR(B{vt + 2}*{last_f}{R["df"]},"N/A")', MONEY),
        ("Enterprise value", f'=IFERROR(B{vt + 1}+B{vt + 3},"N/A")', MONEY),
        ("Less: total debt", "=-B23", MONEY),
        ("Plus: cash & short-term investments", "=B28", MONEY),
        ("Equity value", f'=IFERROR(B{vt + 4}+B{vt + 5}+B{vt + 6},"N/A")', MONEY),
        (f"Value per share ({ccy})", f'=IFERROR(B{vt + 7}/B29,"N/A")', PRICE),
        (f"Value per share ({tccy})", f'=IFERROR(B{vt + 8}*B31,"N/A")', PRICE),
        (f"Current share price ({tccy})", "=B30", PRICE),
        ("Upside / (downside)", f'=IFERROR(B{vt + 9}/B{vt + 10}-1,"N/A")', PCT),
        ("Terminal value as % of enterprise value", f'=IFERROR(B{vt + 3}/B{vt + 4},"N/A")', PCT),
    ]
    for i, (label, formula, fmt) in enumerate(val, start=1):
        put(f"A{vt + i}", label, font=BOLD if "per share" in label or "Upside" in label else None)
        put(f"B{vt + i}", formula, fmt, BOLD if "per share" in label or "Upside" in label else None)

    # ---------- Sensitivity ----------
    st = vt + len(val) + 2
    header(st, f"Sensitivity: value per share ({tccy}). Rows = WACC, columns = terminal growth")
    put(f"A{st + 1}", "WACC ↓ / Terminal growth →", font=BOLD)
    steps = [-2, -1, 0, 1, 2]
    for j, k in enumerate(steps):
        put(f"{col(2 + j)}{st + 1}", f"=$B$13+{k}*0.005", PCT, BOLD)
    for i, k in enumerate(steps):
        r = st + 2 + i
        put(f"A{r}", f"=$B$25+{k}*0.005", PCT, BOLD)
        for j in range(len(steps)):
            gc = f"{col(2 + j)}${st + 1}"
            put(
                f"{col(2 + j)}{r}",
                f'=IF($A{r}<={gc},"N/A",(NPV($A{r},{fcf_range})+{last_f}{R["fcf"]}*(1+{gc})/($A{r}-{gc})/(1+$A{r})^$B$5'
                f"-$B$23+$B$28)/$B$29*$B$31)",
                PRICE,
                BOLD if k == 0 and steps[j] == 0 else None,
            )
    ws[f"A{st + 8}"] = "Centre cell = your base case. Each cell re-runs the valuation with that WACC and growth rate."
    ws[f"A{st + 8}"].font = Font(italic=True, color="64748B")

    ws.column_dimensions["A"].width = 46
    for c in range(2, first_f + n):
        ws.column_dimensions[col(c)].width = 14
    for row in ws.iter_rows(min_row=top, max_row=st + 6, min_col=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="right")
    ws.freeze_panes = "B4"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
